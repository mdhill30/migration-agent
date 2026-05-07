#!/usr/bin/env python3
"""
Generate professional .docx template files for profile_report and DQR.

These templates use placeholder text to show the document structure and styling
that will be applied when real data is rendered via docs_cli.py.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / "tools"))
from docx_styles import (
    Colors, SEVERITY_COLORS,
    create_styled_document, add_cover_page, add_header_footer,
    add_styled_table, add_kv_table, add_section_divider,
    add_callout_box, add_stat_block, add_severity_badge,
)
from docx.shared import Pt


def generate_profile_template():
    """Generate profile_report_template.docx with professional styling."""
    doc = create_styled_document()

    # Cover page
    add_cover_page(doc, "Source Data Profile Report", "{{CUSTOMER}}", "{{VENDOR}}", "{{DATE}}")

    # Header/footer
    add_header_footer(doc, "Source Data Profile Report", "{{CUSTOMER}}")

    # --- Table of Contents placeholder ---
    doc.add_heading("Table of Contents", level=1)
    p = doc.add_paragraph()
    run = p.add_run("{{TOC — auto-generated from headings}}")
    run.font.color.rgb = Colors.MID_GREY
    run.italic = True

    doc.add_page_break()

    # --- Executive Summary ---
    doc.add_heading("Executive Summary", level=1)

    add_stat_block(doc, [
        ("Source Files", "{{N}}"),
        ("Layers", "{{N}}"),
        ("Total Features", "{{N}}"),
        ("CRS", "{{EPSG}}"),
    ])

    p = doc.add_paragraph()
    run = p.add_run("{{EXECUTIVE_SUMMARY — brief description of source data characteristics}}")
    run.font.color.rgb = Colors.MID_GREY
    run.italic = True

    add_section_divider(doc)

    # --- Source Data Summary ---
    doc.add_heading("Source Data Summary", level=1)

    add_styled_table(doc,
        headers=["Layer", "Geometry", "Count", "Key Fields"],
        rows=[
            ["{{LAYER_NAME}}", "{{Point/LineString/Polygon}}", "{{count}}", "{{comma-separated fields}}"],
            ["{{LAYER_NAME}}", "{{Point/LineString/Polygon}}", "{{count}}", "{{comma-separated fields}}"],
        ],
        col_widths=[4.5, 3.5, 2.0, 6.0])

    # --- Layer Inventory ---
    doc.add_heading("Layer Inventory", level=1)
    doc.add_heading("{{SOURCE_FILE_1}}", level=2)

    add_styled_table(doc,
        headers=["Attribute", "Type", "Non-Null", "Unique", "Sample Values"],
        rows=[
            ["{{attr_name}}", "{{varchar(255)}}", "{{count}}", "{{count}}", "{{val1, val2, val3}}"],
        ],
        col_widths=[3.5, 2.5, 2.0, 2.0, 6.0])

    # --- Structural Model ---
    doc.add_heading("Structural Model Assessment", level=1)

    add_callout_box(doc,
        "{{Assessment of the source data model — placement-based vs. explicit containment, "
        "relationship patterns, and implications for NMT mapping.}}",
        style="info")

    # --- Value Domains ---
    doc.add_heading("Value Domains", level=1)
    doc.add_heading("{{LAYER}} — {{ATTRIBUTE}} ({{N}} distinct values)", level=3)

    add_styled_table(doc,
        headers=["Value", "Count", "Interpretation"],
        rows=[
            ["{{value}}", "{{count}}", "{{meaning or 'Unknown'}}"],
        ],
        col_widths=[4.0, 2.5, 9.5])

    # --- Data Quality Issues ---
    doc.add_heading("Data Quality Issues", level=1)

    add_callout_box(doc,
        "{{N}} issue(s) identified during profiling. See DQR document for full details.",
        style="warning")

    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run("{{DQI-001}}: {{Issue description}}")
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run("{{DQI-002}}: {{Issue description}}")

    # --- Open Questions ---
    doc.add_heading("Open Questions", level=1)

    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run("{{Question requiring customer clarification}}")

    # Save
    out_path = Path(__file__).parent.parent / "templates" / "profile_report_template.docx"
    doc.save(out_path)
    print(f"Generated: {out_path}")


def generate_dqr_template():
    """Generate dqr_template.docx with professional styling."""
    doc = create_styled_document()

    # Cover page
    add_cover_page(doc, "Data Quality Review", "{{CUSTOMER}}", "{{VENDOR}}", "{{DATE}}")

    # Header/footer
    add_header_footer(doc, "Data Quality Review", "{{CUSTOMER}}")

    # --- Executive Summary ---
    doc.add_heading("Executive Summary", level=1)

    add_stat_block(doc, [
        ("Total Issues", "{{N}}"),
        ("Open", "{{N}}"),
        ("Resolved", "{{N}}"),
        ("Critical", "{{N}}"),
    ])

    p = doc.add_paragraph()
    run = p.add_run(
        "This document presents the Data Quality Review findings for the {{CUSTOMER}} "
        "data migration project. A total of {{N}} issues were identified during "
        "source data profiling and validation."
    )

    add_callout_box(doc,
        "{{N}} critical/blocker issue(s) require resolution before migration can proceed.",
        style="danger")

    add_section_divider(doc)

    # --- Severity Summary ---
    doc.add_heading("Issue Summary by Severity", level=2)

    add_styled_table(doc,
        headers=["Severity", "Count", "Open"],
        rows=[
            ["Critical", "{{N}}", "{{N}}"],
            ["High", "{{N}}", "{{N}}"],
            ["Medium", "{{N}}", "{{N}}"],
            ["Low", "{{N}}", "{{N}}"],
            ["Info", "{{N}}", "{{N}}"],
        ],
        col_widths=[5.0, 3.0, 3.0],
        accent_col=0)

    # --- Treatment Summary ---
    doc.add_heading("Issue Summary by Treatment", level=2)

    add_styled_table(doc,
        headers=["Treatment", "Count"],
        rows=[
            ["Accept Ignore", "{{N}}"],
            ["Fix In Flight", "{{N}}"],
            ["Needs Decision", "{{N}}"],
        ],
        col_widths=[8.0, 3.0])

    doc.add_page_break()

    # --- Issues Detail ---
    doc.add_heading("Issues Detail", level=1)

    p = doc.add_paragraph()
    run = p.add_run("Each issue is documented with its source, impact, and recommended treatment.")
    run.font.color.rgb = Colors.MID_GREY

    # Sample issue
    heading = doc.add_heading("{{DQI-001}} — {{Issue description}}", level=2)
    for run in heading.runs:
        run.font.color.rgb = Colors.HIGH
    add_severity_badge(heading, "high")

    add_kv_table(doc, [
        ("Category", "{{coded_values}}"),
        ("Source Object", "{{LAYER_NAME}}"),
        ("Source Attribute", "{{ATTRIBUTE_NAME}}"),
        ("Severity", "{{High}}"),
        ("Affected Count", "{{N}}"),
        ("Affected %", "{{N}}%"),
        ("Examples", "{{val1, val2, val3}}"),
        ("Treatment", "{{Fix In Flight}}"),
        ("Status", "{{Open}}"),
        ("Owner", "{{migration_engineer}}"),
    ])

    add_callout_box(doc, "Impact if unresolved: {{description of downstream consequences}}", style="warning")
    add_callout_box(doc, "Resolution: {{description of fix applied}}", style="success")

    # Save
    out_path = Path(__file__).parent.parent / "templates" / "dqr_template.docx"
    doc.save(out_path)
    print(f"Generated: {out_path}")


if __name__ == "__main__":
    generate_profile_template()
    generate_dqr_template()
