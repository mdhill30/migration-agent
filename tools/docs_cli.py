#!/usr/bin/env python3
"""
docs_cli.py — Generate customer-facing documents from migration artefacts.

Usage:
    python tools/docs_cli.py all     --job jobs/<name>
    python tools/docs_cli.py dmdd    --job jobs/<name>
    python tools/docs_cli.py dqr     --job jobs/<name>
    python tools/docs_cli.py profile --job jobs/<name>
"""

import argparse
import os
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re


# =============================================================================
# Helpers
# =============================================================================

def load_yaml(path: Path) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_migration_yaml(job_dir: Path) -> dict:
    path = job_dir / "migration.yaml"
    if path.exists():
        return load_yaml(path)
    return {}


def ensure_output_dir(job_dir: Path) -> Path:
    out = job_dir / "output"
    out.mkdir(parents=True, exist_ok=True)
    return out


# =============================================================================
# DMDD → Excel
# =============================================================================

# Style constants
HEADER_FONT = Font(name="Calibri", size=14, bold=True)
SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True)
DATA_FONT = Font(name="Calibri", size=11)
HEADER_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _set_col_widths(ws, widths: dict):
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


def _write_header_row(ws, row: int, headers: list, start_col: int = 1):
    for i, h in enumerate(headers, start=start_col):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = SUBHEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_data_row(ws, row: int, values: list, start_col: int = 1):
    for i, v in enumerate(values, start=start_col):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def generate_dmdd_excel(job_dir: Path):
    """Generate DMDD Excel from dmdd.yaml."""
    dmdd_path = job_dir / "dmdd.yaml"
    if not dmdd_path.exists():
        print(f"ERROR: {dmdd_path} not found")
        return None

    dmdd = load_yaml(dmdd_path)
    migration = load_migration_yaml(job_dir)
    summary = dmdd.get("summary", {})

    wb = Workbook()

    # --- Summary sheet ---
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value="Vendor").font = SUBHEADER_FONT
    ws.cell(row=1, column=2, value=summary.get("vendor", "")).font = DATA_FONT
    ws.cell(row=2, column=1, value="Customer").font = SUBHEADER_FONT
    ws.cell(row=2, column=2, value=summary.get("customer", "")).font = DATA_FONT
    ws.cell(row=3, column=1, value="Date").font = SUBHEADER_FONT
    ws.cell(row=3, column=2, value=summary.get("date", "")).font = DATA_FONT
    ws.cell(row=5, column=1, value="Instructions").font = SUBHEADER_FONT
    for i, instr in enumerate(summary.get("instructions", []), start=6):
        ws.cell(row=i, column=1, value=instr).font = DATA_FONT

    status_row = 6 + len(summary.get("instructions", []))  + 1
    ws.cell(row=status_row, column=1, value="Status Values").font = SUBHEADER_FONT
    for i, sv in enumerate(summary.get("status_values", []), start=status_row + 1):
        ws.cell(row=i, column=1, value=sv).font = DATA_FONT

    _set_col_widths(ws, {"A": 20, "B": 60})

    # --- ObjectInventory sheet ---
    ws = wb.create_sheet("ObjectInventory")
    headers = ["Feature", "Count", "Include in Mapping?", "Geometry", "Notes"]
    _write_header_row(ws, 1, headers)
    for i, obj in enumerate(dmdd.get("object_inventory", []), start=2):
        _write_data_row(ws, i, [
            obj.get("feature", ""),
            obj.get("count", ""),
            obj.get("include_in_mapping", ""),
            obj.get("geometry", ""),
            obj.get("notes", ""),
        ])
    _set_col_widths(ws, {"A": 35, "B": 12, "C": 20, "D": 15, "E": 60})

    # --- AttributeInventory sheet ---
    ws = wb.create_sheet("AttributeInventory")
    headers = ["Feature", "Attribute", "Data Type", "Count NULL", "Count Unique", "Include In Mapping?"]
    _write_header_row(ws, 1, headers)
    for i, attr in enumerate(dmdd.get("attribute_inventory", []), start=2):
        _write_data_row(ws, i, [
            attr.get("feature", ""),
            attr.get("attribute", ""),
            attr.get("data_type", ""),
            attr.get("count_null", ""),
            attr.get("count_unique", ""),
            attr.get("include_in_mapping", ""),
        ])
    _set_col_widths(ws, {"A": 25, "B": 25, "C": 18, "D": 12, "E": 14, "F": 20})

    # --- ObjectMapping sheet ---
    ws = wb.create_sheet("ObjectMapping")
    # Section headers (row 1)
    ws.cell(row=1, column=1, value="Tracking").font = HEADER_FONT
    ws.cell(row=1, column=6, value="IQGeo Features").font = HEADER_FONT
    ws.cell(row=1, column=16, value="Source Mapping").font = HEADER_FONT

    # Column headers (row 2)
    obj_map_headers = [
        "Update Person", "Update Date", "Review Comments", "Status", "Jira Ticket",  # A-E (Tracking)
        "Feature", "Display Name", "Internal Name", "Editable", "Layer",             # F-J (IQGeo)
        "Specifications", "Feature Type", "Function", "Housing", "Style",            # K-O (IQGeo cont.)
        "System", "Table/Layer", "Transformations", "Internal Name",                 # P-S (Source)
    ]
    _write_header_row(ws, 2, obj_map_headers)

    for i, mapping in enumerate(dmdd.get("object_mapping", []), start=3):
        tracking = mapping.get("tracking", {})
        iqgeo = mapping.get("iqgeo_features", {})
        source = mapping.get("source_mapping", {})
        _write_data_row(ws, i, [
            tracking.get("update_person", ""),
            tracking.get("update_date", ""),
            tracking.get("review_comments", ""),
            tracking.get("status", ""),
            tracking.get("jira_ticket", ""),
            iqgeo.get("feature", ""),
            iqgeo.get("display_name", ""),
            iqgeo.get("internal_name", ""),
            iqgeo.get("editable", ""),
            iqgeo.get("layer", ""),
            iqgeo.get("specifications", ""),
            iqgeo.get("feature_type", ""),
            iqgeo.get("function", ""),
            iqgeo.get("housing", ""),
            iqgeo.get("style", ""),
            source.get("system", ""),
            source.get("table_layer", ""),
            source.get("transformations", ""),
            source.get("internal_name", ""),
        ])
    _set_col_widths(ws, {
        "A": 14, "B": 12, "C": 25, "D": 22, "E": 12,
        "F": 22, "G": 22, "H": 22, "I": 10, "J": 12,
        "K": 18, "L": 14, "M": 14, "N": 14, "O": 12,
        "P": 14, "Q": 30, "R": 35, "S": 20,
    })

    # --- AttributeMapping sheet ---
    ws = wb.create_sheet("AttributeMapping")
    ws.cell(row=1, column=1, value="Tracking").font = HEADER_FONT
    ws.cell(row=1, column=5, value="IQGeo Internal").font = HEADER_FONT
    ws.cell(row=1, column=7, value="IQGeo Fields").font = HEADER_FONT
    ws.cell(row=1, column=16, value="Source Mapping").font = HEADER_FONT

    attr_map_headers = [
        "Update Person", "Update Date", "Review Comments", "Status",              # A-D
        "IQGeo Version", "Jira Ticket",                                           # E-F
        "Table", "Table Type", "Display Name", "Internal Name", "Data Type",      # G-K
        "Default Value", "Pick List", "Visible", "Other",                         # L-O
        "System", "Table/Layer", "Attribute", "Transformations",                  # P-S
        "Internal Name", "Data Type", "Length", "Domain", "Sample Value",         # T-X
    ]
    _write_header_row(ws, 2, attr_map_headers)

    for i, mapping in enumerate(dmdd.get("attribute_mapping", []), start=3):
        tracking = mapping.get("tracking", {})
        iqgeo_int = mapping.get("iqgeo_internal", {})
        iqgeo_f = mapping.get("iqgeo_fields", {})
        source = mapping.get("source_mapping", {})
        _write_data_row(ws, i, [
            tracking.get("update_person", ""),
            tracking.get("update_date", ""),
            tracking.get("review_comments", ""),
            tracking.get("status", ""),
            iqgeo_int.get("iqgeo_version_tested", ""),
            iqgeo_int.get("jira_ticket", ""),
            iqgeo_f.get("table", ""),
            iqgeo_f.get("table_type", ""),
            iqgeo_f.get("display_name", ""),
            iqgeo_f.get("internal_name", ""),
            iqgeo_f.get("data_type", ""),
            iqgeo_f.get("default_value", ""),
            iqgeo_f.get("pick_list", ""),
            iqgeo_f.get("visible", ""),
            iqgeo_f.get("other", ""),
            source.get("system", ""),
            source.get("table_layer", ""),
            source.get("attribute", ""),
            source.get("transformations", ""),
            source.get("internal_name", ""),
            source.get("data_type", ""),
            source.get("length", ""),
            source.get("domain", ""),
            source.get("sample_value", ""),
        ])
    _set_col_widths(ws, {
        "A": 14, "B": 12, "C": 25, "D": 22,
        "E": 14, "F": 12,
        "G": 18, "H": 12, "I": 20, "J": 20, "K": 16,
        "L": 14, "M": 14, "N": 10, "O": 20,
        "P": 14, "Q": 20, "R": 20, "S": 30,
        "T": 18, "U": 14, "V": 8, "W": 16, "X": 20,
    })

    # --- Value mapping sheets (AttrVal-*) ---
    for key, val in dmdd.items():
        if key.startswith("value_mapping_"):
            domain_name = key.replace("value_mapping_", "")
            sheet_name = f"AttrVal-{domain_name}"[:31]  # Excel 31-char limit
            ws = wb.create_sheet(sheet_name)

            ws.cell(row=1, column=2, value="Source System").font = SUBHEADER_FONT
            ws.cell(row=1, column=4, value="Target System").font = SUBHEADER_FONT

            ws.cell(row=2, column=1, value="Source Value").font = SUBHEADER_FONT
            ws.cell(row=2, column=2, value="Count").font = SUBHEADER_FONT
            ws.cell(row=2, column=3, value="").font = SUBHEADER_FONT
            ws.cell(row=2, column=4, value="Target Value").font = SUBHEADER_FONT
            ws.cell(row=2, column=5, value="Notes").font = SUBHEADER_FONT

            _write_header_row(ws, 2, ["Source Value", "Count", "", "Target Value", "Notes"])

            if isinstance(val, list):
                for j, entry in enumerate(val, start=3):
                    _write_data_row(ws, j, [
                        entry.get("source_value", ""),
                        entry.get("count", ""),
                        "",
                        entry.get("target_value", ""),
                        entry.get("notes", ""),
                    ])
            _set_col_widths(ws, {"A": 25, "B": 10, "C": 3, "D": 25, "E": 40})

    # Save
    out_dir = ensure_output_dir(job_dir)
    out_path = out_dir / "dmdd.xlsx"
    wb.save(out_path)
    print(f"Generated: {out_path}")
    return out_path


# =============================================================================
# DQR → Word
# =============================================================================

SEVERITY_COLORS = {
    "blocker": RGBColor(0xC0, 0x00, 0x00),
    "high": RGBColor(0xE0, 0x60, 0x00),
    "medium": RGBColor(0xBF, 0x8F, 0x00),
    "low": RGBColor(0x2E, 0x74, 0xB5),
    "info": RGBColor(0x70, 0x70, 0x70),
}


def generate_dqr_word(job_dir: Path):
    """Generate DQR Word document from dqr.yaml."""
    dqr_path = job_dir / "dqr.yaml"
    if not dqr_path.exists():
        print(f"ERROR: {dqr_path} not found")
        return None

    dqr = load_yaml(dqr_path)
    migration = load_migration_yaml(job_dir)
    issues = dqr.get("issues", [])

    # Count by severity
    severity_counts = {}
    for issue in issues:
        sev = issue.get("severity", "info")
        severity_counts[sev] = severity_counts.get(sev, 0) + 1

    # Get metadata
    summary = dqr.get("summary", {}) or {}
    customer = summary.get("customer") or migration.get("customer", "")
    vendor = summary.get("vendor") or migration.get("vendor", "IQGeo")
    date = summary.get("date") or migration.get("date", "")

    doc = Document()

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title page
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Data Quality Review")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(customer)
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Prepared by: {vendor}").font.size = Pt(12)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta2.add_run(f"Date: {date}").font.size = Pt(12)

    doc.add_page_break()

    # Executive Summary
    doc.add_heading("Executive Summary", level=1)
    total = len(issues)
    open_count = sum(1 for i in issues if i.get("status", "open") == "open")
    doc.add_paragraph(
        f"This document presents the Data Quality Review findings for the {customer} "
        f"data migration project. A total of {total} issues were identified, "
        f"of which {open_count} remain open."
    )

    # Summary table
    doc.add_heading("Issue Summary by Severity", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Severity"
    hdr[1].text = "Count"
    hdr[2].text = "Open"

    for sev in ["blocker", "high", "medium", "low", "info"]:
        count = severity_counts.get(sev, 0)
        open_c = sum(1 for i in issues if i.get("severity") == sev and i.get("status", "open") == "open")
        row = table.add_row().cells
        row[0].text = sev.capitalize()
        row[1].text = str(count)
        row[2].text = str(open_c)

    # Treatment summary
    doc.add_heading("Issue Summary by Treatment", level=2)
    treatment_counts = {}
    for issue in issues:
        t = issue.get("treatment", "unknown")
        treatment_counts[t] = treatment_counts.get(t, 0) + 1

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Treatment"
    hdr[1].text = "Count"
    for t, c in sorted(treatment_counts.items()):
        row = table.add_row().cells
        row[0].text = t.replace("_", " ").title()
        row[1].text = str(c)

    doc.add_page_break()

    # Issues Detail
    doc.add_heading("Issues Detail", level=1)

    for issue in issues:
        issue_id = issue.get("id", "")
        check = issue.get("check", "")
        severity = issue.get("severity", "info")

        heading = doc.add_heading(f"{issue_id} — {check}", level=2)
        # Color the heading by severity
        for run in heading.runs:
            color = SEVERITY_COLORS.get(severity)
            if color:
                run.font.color.rgb = color

        # Issue details table
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"

        fields = [
            ("Category", issue.get("category", "")),
            ("Source Object", issue.get("source_object", "")),
            ("Source Attribute", issue.get("source_attribute", "")),
            ("Severity", severity.capitalize()),
            ("Affected Count", str(issue.get("affected_count", ""))),
            ("Affected %", f"{issue.get('affected_percent', '')}%"),
            ("Examples", issue.get("examples", "")),
            ("Treatment", issue.get("treatment", "").replace("_", " ").title()),
            ("Target Area", issue.get("target_area", "")),
            ("Status", issue.get("status", "")),
            ("Owner", issue.get("owner", "")),
        ]

        for label, value in fields:
            row = table.add_row().cells
            row[0].text = label
            p = row[0].paragraphs[0]
            p.runs[0].bold = True
            row[1].text = str(value) if value else ""

        # Impact
        impact = issue.get("impact_if_unresolved", "")
        if impact:
            p = doc.add_paragraph()
            p.add_run("Impact if unresolved: ").bold = True
            p.add_run(impact)

        # Resolution
        resolution = issue.get("resolution", "")
        if resolution:
            p = doc.add_paragraph()
            p.add_run("Resolution: ").bold = True
            p.add_run(resolution)

        doc.add_paragraph()  # spacer

    # Save
    out_dir = ensure_output_dir(job_dir)
    out_path = out_dir / "dqr.docx"
    doc.save(out_path)
    print(f"Generated: {out_path}")
    return out_path


# =============================================================================
# Profile Report → Word
# =============================================================================

def _parse_markdown_table(lines: list) -> list:
    """Parse markdown table lines into list of rows (each row is list of cells)."""
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        # Skip separator rows
        if re.match(r"^\|[\s\-:|]+\|$", line):
            continue
        cells = [c.strip() for c in line.split("|")[1:-1]]
        rows.append(cells)
    return rows


def generate_profile_word(job_dir: Path):
    """Generate profile report Word document from profile_report.md."""
    report_path = job_dir / "profile_report.md"
    if not report_path.exists():
        print(f"ERROR: {report_path} not found")
        return None

    migration = load_migration_yaml(job_dir)
    customer = migration.get("customer", "")
    vendor = migration.get("vendor", "IQGeo")
    date = migration.get("date", "")

    with open(report_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title page
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Source Data Profile Report")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(customer)
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Prepared by: {vendor}").font.size = Pt(12)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta2.add_run(f"Date: {date}").font.size = Pt(12)

    doc.add_page_break()

    # Parse markdown and convert to Word
    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        # Headings
        if line.startswith("# "):
            # Skip the first H1 (it's the title, already on title page)
            if i == 0 or (i > 0 and not any(l.startswith("# ") for l in lines[:i])):
                i += 1
                continue
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        # Horizontal rule
        if line.strip() == "---":
            i += 1
            continue

        # Tables
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = _parse_markdown_table(table_lines)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r_idx, row_data in enumerate(rows):
                    for c_idx, cell_val in enumerate(row_data):
                        if c_idx < len(table.columns):
                            table.cell(r_idx, c_idx).text = cell_val
                            if r_idx == 0:
                                p = table.cell(r_idx, c_idx).paragraphs[0]
                                if p.runs:
                                    p.runs[0].bold = True
                doc.add_paragraph()  # spacer after table
            continue

        # Code blocks
        if line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        # Bold text, bullet points, regular paragraphs
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            text = line.strip()[2:]
            p = doc.add_paragraph(style="List Bullet")
            _add_formatted_text(p, text)
            i += 1
            continue

        if line.strip():
            p = doc.add_paragraph()
            _add_formatted_text(p, line.strip())
            i += 1
            continue

        i += 1

    # Save
    out_dir = ensure_output_dir(job_dir)
    out_path = out_dir / "profile_report.docx"
    doc.save(out_path)
    print(f"Generated: {out_path}")
    return out_path


def _add_formatted_text(paragraph, text: str):
    """Add text to paragraph with basic markdown formatting (bold, italic, code)."""
    # Process **bold**, *italic*, `code` inline formatting
    pattern = r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|([^*`]+))"
    for match in re.finditer(pattern, text):
        if match.group(2):  # bold
            run = paragraph.add_run(match.group(2))
            run.bold = True
        elif match.group(3):  # italic
            run = paragraph.add_run(match.group(3))
            run.italic = True
        elif match.group(4):  # code
            run = paragraph.add_run(match.group(4))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif match.group(5):  # plain text
            paragraph.add_run(match.group(5))


# =============================================================================
# CLI
# =============================================================================

def main():
    parser = argparse.ArgumentParser(description="Generate customer-facing documents from migration artefacts")
    parser.add_argument("command", choices=["all", "dmdd", "dqr", "profile"],
                        help="Which document(s) to generate")
    parser.add_argument("--job", required=True, help="Path to job directory")

    args = parser.parse_args()
    job_dir = Path(args.job)

    if not job_dir.exists():
        print(f"ERROR: Job directory not found: {job_dir}")
        sys.exit(1)

    if args.command in ("all", "dmdd"):
        generate_dmdd_excel(job_dir)
    if args.command in ("all", "dqr"):
        generate_dqr_word(job_dir)
    if args.command in ("all", "profile"):
        generate_profile_word(job_dir)


if __name__ == "__main__":
    main()
